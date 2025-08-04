"""
Type A Training Routes - ECBC Compliance Training
===============================================

Routes and logic for Type A (ECBC Compliance) training reports.
This contains the exact same logic as your current app_clean.py.
"""
from flask import Blueprint, render_template, request, redirect, url_for
import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Cm

# Import existing modules (no changes needed)
from modules.document_utils import find_and_replace_text, find_and_replace_image, save_uploaded_file
from modules.image_processing import insert_gallery_table, get_annexure_images_and_captions, insert_annexure_images
from modules.form_processing import process_form_data, process_gallery_images
from config import Config

# Create Type A blueprint
type_a_bp = Blueprint('type_a', __name__)

@type_a_bp.route('/')
def form():
    """Type A Training form page."""
    print("📋 Type A form page accessed")
    return render_template('type_a/form.html')

@type_a_bp.route('/generate', methods=['POST'])
def generate_report():
    """Generate Type A training report - EXACT logic from your app_clean.py."""
    try:
        print("🎯 Type A generate route accessed")
        print(f"📝 Form data received: {dict(request.form)}")
        
        # Get the selected template number
        selected_template = request.form.get('selected_template', '1')
        organization = request.form.get('cell_name', 'Unknown')
        
        print(f"DEBUG: Selected template: {selected_template}")
        print(f"DEBUG: Organization: {organization}")
        print(f"DEBUG: All form data keys: {list(request.form.keys())}")
        
        # Map template numbers to file names
        template_mapping = {
            '1': 'templates/type_a/word_templates/word_template_1.docx',  # RRECL
            '2': 'templates/type_a/word_templates/word_template_2.docx',  # GEDA
            '3': 'templates/type_a/word_templates/word_template_3.docx',  # HAREDA
            '4': 'templates/type_a/word_templates/word_template_4.docx',  # UREDA
            '5': 'templates/type_a/word_templates/word_template_5.docx'   # SDA Odisha
        }
        
        # Get the template file name
        template_file = template_mapping.get(selected_template, 'templates/type_a/word_templates/word_template_1.docx')
        
        print(f"DEBUG: Using template file: {template_file}")
        
        # Check if template file exists (use absolute path)
        if not os.path.exists(os.path.abspath(template_file)):
            return f"Error: Template file '{template_file}' not found at {os.path.abspath(template_file)}. Please ensure all template files are present.", 400
        
        # Load the Word template
        doc = Document(template_file)

        # Process form data
        text_replacements = process_form_data(request)
        
        # Apply text replacements
        for placeholder, value in text_replacements.items():
            if value:
                find_and_replace_text(doc, placeholder, value)

        # Process gallery images
        gallery_images_clean, gallery_captions_clean = process_gallery_images(request)
        
        # Insert the gallery table if images exist with 2×3 layout (6 images per page)
        if gallery_images_clean:
            insert_gallery_table(doc, gallery_images_clean, gallery_captions_clean, 
                                images_per_row=2, image_width=Cm(8.13))

        # Process annexure images with improved dimensions
        annexure_placeholders = [
            ('annexure1', '{{ANNEXURE1_TABLE}}'),
            ('annexure2', '{{ANNEXURE2_TABLE}}'),
            ('annexure3', '{{ANNEXURE3_TABLE}}'),
            ('annexure4', '{{ANNEXURE4_TABLE}}'),
            ('annexure5', '{{ANNEXURE5_TABLE}}'),
        ]

        for i, (prefix, placeholder) in enumerate(annexure_placeholders):
            images, captions = get_annexure_images_and_captions(prefix, request)
            if images:
                # Only skip page break for the last annexure (annexure5)
                is_last_annexure = (i == len(annexure_placeholders) - 1)
                insert_annexure_images(doc, images, captions, placeholder, 
                                     image_width=Cm(15), image_height=Cm(20),
                                     add_final_page_break=not is_last_annexure)

        # Generate output filename
        event_date = request.form.get('event_date', '').replace('-', '')
        cell_name = request.form.get('cell_name', '').replace(' ', '_')
        filename = f"TypeA_{event_date}_{cell_name}_report.docx"
        output_path = os.path.join(Config.OUTPUT_FOLDER, filename)
        
        # Save the document
        doc.save(output_path)
        
        # After successful generation, redirect to success page
        return redirect(url_for('type_a.success', filename=os.path.basename(output_path)))
        
    except Exception as e:
        print(f"❌ Error: {str(e)}")
        return render_template('error.html', error=str(e))

@type_a_bp.route('/success')
def success():
    """Display success page with download option."""
    filename = request.args.get('filename')
    if not filename:
        return redirect(url_for('type_a.form'))
    
    download_url = url_for('download_file', filename=filename)
    return render_template('type_a/success.html', filename=filename, download_url=download_url)
