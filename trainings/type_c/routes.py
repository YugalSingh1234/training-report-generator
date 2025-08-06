""""
Type C Training Routes - Professional Development Training
========================================================

Routes and logic for Type C (Professional Development) training reports.
Handles gallery images and 6 annexure documents with table insertion.
"""
from flask import Blueprint, render_template, request, redirect, url_for, current_app
import os
import sys
from datetime import datetime
from docx import Document
from docx.shared import Inches, Cm

# Import existing modules (using sys.path to resolve from parent directory)
sys.path.append(os.path.join(os.path.dirname(__file__), '..', '..'))
from modules.document_utils import find_and_replace_text, find_and_replace_image, save_uploaded_file
from modules.image_processing import insert_gallery_table, get_annexure_images_and_captions, insert_annexure_images
from modules.form_processing import process_form_data, process_gallery_images

# Create Type C blueprint
type_c_bp = Blueprint('type_c', __name__, url_prefix='/type-c')

def process_type_c_form_data(request):
    """Process Type C specific form data with the exact placeholders provided."""
    # Get basic form data first
    basic_data = process_form_data(request)
    
    # Handle date range formatting for Type C (start_date and end_date)
    start_date = request.form.get('start_date', '')
    end_date = request.form.get('end_date', '')
    
    if start_date and end_date:
        try:
            # Parse dates and format as "6th & 6th August 2025" or "5th & 6th August 2025"
            from datetime import datetime
            start_dt = datetime.strptime(start_date, '%Y-%m-%d')
            end_dt = datetime.strptime(end_date, '%Y-%m-%d')
            
            # Format with ordinal suffixes
            def get_ordinal(day):
                if 10 <= day % 100 <= 20:
                    suffix = 'th'
                else:
                    suffix = {1: 'st', 2: 'nd', 3: 'rd'}.get(day % 10, 'th')
                return f"{day}{suffix}"
            
            start_day = get_ordinal(start_dt.day)
            end_day = get_ordinal(end_dt.day)
            month_year = start_dt.strftime('%B %Y')
            
            # Handle same date case (like "6th & 6th August 2025")
            if start_dt.date() == end_dt.date():
                formatted_date = f"{start_day} & {end_day} {month_year}"
            else:
                # Different dates (like "5th & 6th August 2025")
                if start_dt.month == end_dt.month and start_dt.year == end_dt.year:
                    formatted_date = f"{start_day} & {end_day} {month_year}"
                else:
                    # Different months/years
                    start_month_year = start_dt.strftime('%B %Y')
                    end_month_year = end_dt.strftime('%B %Y')
                    formatted_date = f"{start_day} {start_month_year} & {end_day} {end_month_year}"
            
            basic_data['{{EVENT_DATE}}'] = formatted_date
        except:
            # Fallback to single date if parsing fails
            basic_data['{{EVENT_DATE}}'] = request.form.get('start_date', '')
    
    # Process address from 3 lines into 2 lines format
    address_line1 = request.form.get('address_line1', '')
    address_line2 = request.form.get('address_line2', '')
    address_line3 = request.form.get('address_line3', '')
    
    # Combine address into 2 lines for output
    if address_line1 and address_line2 and address_line3:
        formatted_address = f"{address_line1}\n{address_line2}, {address_line3}"
    elif address_line1 and address_line2:
        formatted_address = f"{address_line1}\n{address_line2}"
    else:
        formatted_address = address_line1 or request.form.get('address', '')
    
    # Process dynamic officials
    senior_officials = []
    i = 1
    while True:
        official_name = request.form.get(f'senior_official_{i}', '')
        official_designation = request.form.get(f'senior_official_designation_{i}', '')
        if not official_name:
            break
        if official_designation:
            senior_officials.append(f"{official_name}, {official_designation}")
        else:
            senior_officials.append(official_name)
        i += 1
    
    senior_officials_text = '; '.join(senior_officials) if senior_officials else ''
    
    # Process dynamic trainers
    trainers = []
    i = 1
    while True:
        trainer_name = request.form.get(f'trainer_name_{i}', '')
        trainer_designation = request.form.get(f'trainer_designation_{i}', '')
        if not trainer_name:
            break
        if trainer_designation:
            trainers.append(f"{trainer_name}, {trainer_designation}")
        else:
            trainers.append(trainer_name)
        i += 1
    
    # Get first two trainers for specific placeholders
    trainer_1 = trainers[0] if len(trainers) > 0 else ''
    trainer_2 = trainers[1] if len(trainers) > 1 else ''
    
    # Chief Guest with designation
    chief_guest_name = request.form.get('chief_guest_name', '')
    chief_guest_designation = request.form.get('chief_guest_designation', '')
    chief_guest_full = f"{chief_guest_name}, {chief_guest_designation}" if chief_guest_name and chief_guest_designation else chief_guest_name
    
    # Type C specific placeholders as provided by user
    type_c_data = {
        '{{EVENT_DATE}}': basic_data.get('{{EVENT_DATE}}', request.form.get('start_date', '')),
        '{{ADDRESS}}': formatted_address,
        '{{Submitted_to}}': request.form.get('submitted_to', ''),
        '{{Submitted_by}}': request.form.get('submitted_by', ''),
        '{{Senior_Official}}': senior_officials_text,
        '{{Chief_Guest_Name}}': chief_guest_full,
        '{{Trainer_Name_1}}': trainer_1,
        '{{Trainer_Name_2}}': trainer_2,
        '{{ Participant_Department }}': request.form.get('participant_department', ''),  # With spaces
        '{{Participant_Department}}': request.form.get('participant_department', ''),    # Without spaces
        '{{ Participant_No. }}': request.form.get('participant_no', ''),                 # With spaces
        '{{Participant_No.}}': request.form.get('participant_no', ''),                   # Without spaces
        '{{CELL_NAME}}': request.form.get('cell_name', ''),
    }
    
    # Merge with basic data (Type C data takes precedence)
    basic_data.update(type_c_data)
    return basic_data

@type_c_bp.route('/')
def form():
    """Type C Training form page."""
    print("📋 Type C form page accessed")
    return render_template('type_c/form.html')

@type_c_bp.route('/generate', methods=['POST'])
def generate_report():
    """Generate Type C training report with gallery and 6 annexures."""
    try:
        print("🎯 Type C generate route accessed")
        print(f"📝 Form data received: {dict(request.form)}")
        
        # Get the selected template number
        selected_template = request.form.get('selected_template', '1')
        organization = request.form.get('cell_name', 'Unknown')
        
        print(f"DEBUG: Selected template: {selected_template}")
        print(f"DEBUG: Organization: {organization}")
        print(f"DEBUG: All form data keys: {list(request.form.keys())}")
        
        # Map template numbers to Type C file names (using absolute paths)
        base_path = current_app.root_path
        template_mapping = {
            '1': os.path.join(base_path, 'templates', 'type_c', 'word_template_1.docx'),
            '2': os.path.join(base_path, 'templates', 'type_c', 'word_template_2.docx'),
            '3': os.path.join(base_path, 'templates', 'type_c', 'word_template_3.docx'),
            '4': os.path.join(base_path, 'templates', 'type_c', 'word_template_4.docx'),
            '5': os.path.join(base_path, 'templates', 'type_c', 'word_template_5.docx')
        }
        
        # Get the template file name
        template_file = template_mapping.get(selected_template, os.path.join(base_path, 'templates', 'type_c', 'word_template_1.docx'))
        
        print(f"DEBUG: Using template file: {template_file}")
        
        # Check if template file exists (use absolute path)
        if not os.path.exists(os.path.abspath(template_file)):
            return f"Error: Template file '{template_file}' not found at {os.path.abspath(template_file)}. Please ensure all template files are present.", 400
        
        # Load the Word template
        doc = Document(template_file)

        # Process Type C specific form data
        text_replacements = process_type_c_form_data(request)
        
        # Apply text replacements
        for placeholder, value in text_replacements.items():
            if value:
                find_and_replace_text(doc, placeholder, value)

        # Process gallery images
        gallery_images_clean, gallery_captions_clean = process_gallery_images(request)
        
        print(f"🖼️ Gallery images processed: {len(gallery_images_clean)} images")
        print(f"📝 Gallery captions: {gallery_captions_clean}")
        
        # Insert the gallery table if images exist with 2×3 layout (6 images per page)
        if gallery_images_clean:
            print("🏗️ Inserting gallery table...")
            insert_gallery_table(doc, gallery_images_clean, gallery_captions_clean, 
                                images_per_row=2, image_width=Cm(8.13))
            print("✅ Gallery table inserted")
        else:
            print("⚠️ No gallery images to insert")
            # Remove the {{GALLERY_TABLE}} placeholder even if no images
            find_and_replace_text(doc, '{{GALLERY_TABLE}}', 'No gallery images uploaded')

        # Process annexure images for Type C (6 annexures)
        annexure_placeholders = [
            ('annexure1', '{{ANNEXURE1_TABLE}}'),  # Annexure-I (Flyer of the Training)
            ('annexure2', '{{ANNEXURE2_TABLE}}'),  # Annexure-II (Attendance Sheet)
            ('annexure3', '{{ANNEXURE3_TABLE}}'),  # Annexure-III (Feedback Form)
            ('annexure4', '{{ANNEXURE4_TABLE}}'),  # Annexure-IV (Registration Form)
            ('annexure5', '{{ANNEXURE5_TABLE}}'),  # Annexure-V (Registration Form continued)
            ('annexure6', '{{ANNEXURE6_TABLE}}'),  # Annexure-VI (Brochure)
        ]

        for i, (prefix, placeholder) in enumerate(annexure_placeholders):
            images, captions = get_annexure_images_and_captions(prefix, request)
            if images:
                print(f"📎 Processing {placeholder} with {len(images)} images")
                # Only skip page break for the last annexure (annexure6)
                is_last_annexure = (i == len(annexure_placeholders) - 1)
                insert_annexure_images(doc, images, captions, placeholder, 
                                     image_width=Cm(15), image_height=Cm(20),
                                     add_final_page_break=not is_last_annexure)
            else:
                print(f"⚠️ No images for {placeholder}, removing placeholder")
                # Remove placeholder if no images
                find_and_replace_text(doc, placeholder, 'No images uploaded for this annexure')

        # Generate output filename for Type C
        start_date = request.form.get('start_date', request.form.get('event_date', '')).replace('-', '')
        cell_name = request.form.get('cell_name', '').replace(' ', '_')
        filename = f"TypeC_{start_date}_{cell_name}_report.docx"
        output_path = os.path.join(current_app.root_path, 'output', filename)
        
        # Save the document
        doc.save(output_path)
        
        # After successful generation, redirect to success page
        return redirect(url_for('type_c.success', filename=os.path.basename(output_path)))
        
    except Exception as e:
        print(f"❌ Error: {str(e)}")
        return render_template('error.html', error=str(e))

@type_c_bp.route('/success')
def success():
    """Display success page with download option."""
    filename = request.args.get('filename')
    if not filename:
        return redirect(url_for('type_c.form'))
    
    download_url = url_for('download_file', filename=filename)
    return render_template('type_c/success.html', filename=filename, download_url=download_url)
